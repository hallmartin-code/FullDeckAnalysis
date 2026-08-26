"""Generate synthetic pitch decks for testing.

    python tests/make_sample_deck.py            # writes input/sample_deck.pptx and .pdf

The content is fictional and deliberately imperfect: it has a top-down-only TAM, a
"no direct competitors" claim, one very dense slide and no risk slide, so the review
pipeline has real material to criticise.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

SLIDES = [
    ("Nimbus Health", [
        "AI-native care coordination for community hospitals",
        "Seed round | Austin, TX | Confidential",
    ]),
    ("The Problem", [
        "Community hospitals lose $4.1M per year to avoidable readmissions.",
        "Care coordinators manage 60+ patients each using fax, phone and spreadsheets.",
        "34% of discharge plans are never followed up within 7 days.",
        "Source: AHA Annual Survey 2024; internal interviews with 41 coordinators.",
    ]),
    ("Our Solution", [
        "Nimbus is a care coordination platform that auto-generates discharge follow-up",
        "plans from the EHR and routes them to the right coordinator.",
        "Integrates with Epic and Cerner via FHIR. Deploys in 6 weeks.",
    ]),
    ("Product", [
        "Three modules: Discharge Copilot, Coordinator Queue, Outcomes Dashboard",
        "Live in production at 7 hospitals since March 2024",
        "SOC 2 Type II certified; HIPAA compliant",
    ]),
    ("Why Now", [
        "CMS readmission penalties expanded to 6 new DRGs in 2025.",
        "FHIR API mandate makes EHR integration viable for the first time.",
    ]),
    ("Market Opportunity", [
        "TAM: $12B US care coordination software market",
        "The global healthcare IT industry is worth $390B and growing 15% annually.",
    ]),
    ("Traction", [
        "ARR $1.4M, up from $310K twelve months ago",
        "7 paying hospital customers; 3 more in contracting",
        "Net revenue retention 128%; logo churn 0 since launch",
        "Average contract value $198K, 14 month average sales cycle",
        "Pilot at Ascension expanded to 4 additional sites in Q1",
    ]),
    ("Business Model", [
        "Annual SaaS subscription priced per staffed bed.",
        "Average contract value $198K. Gross margin 78%.",
        "Land with one service line, expand across the system.",
    ]),
    ("Go-To-Market", [
        "Direct enterprise sales to VP of Care Management.",
        "Channel partnership with a regional GPO covering 240 hospitals.",
        "CAC $61K, payback 11 months.",
    ]),
    ("Competition", [
        "We have no direct competitors in AI-native discharge coordination.",
        "Legacy vendors: Epic Care Everywhere, Cerner HealtheIntent, Innovaccer.",
        "These are analytics tools, not workflow tools.",
    ]),
    ("Financial Detail and Operating Plan", [
        "The operating plan assumes we grow from $1.4M ARR to $4.2M ARR over the next "
        "eighteen months, driven primarily by expansion within existing health systems "
        "and secondarily by new logo acquisition through the GPO channel partnership. "
        "We model a 128% net revenue retention rate held flat, which we believe is "
        "conservative given the expansion pattern observed at Ascension and two other "
        "systems where we started in a single service line and expanded to three or "
        "more within nine months. Sales headcount grows from two account executives to "
        "six over the plan period, with each fully ramped rep expected to close "
        "approximately $900K of new annual contract value per year at a fully loaded "
        "cost of $310K including quota-carrying support. Gross margin is modeled at 78% "
        "declining slightly to 76% as implementation services grow as a share of "
        "revenue. Monthly burn is currently $340K and rises to $520K at peak hiring "
        "before declining as revenue scales. We reach cash flow breakeven in month 31 "
        "under the base case and month 24 under the upside case.",
    ]),
    ("Team", [
        "Dr. Amara Okafor - CEO. Former VP Care Management, Baylor Scott & White.",
        "Ravi Chandrasekaran - CTO. Engineering lead at Epic for 8 years.",
        "Marta Lindqvist - Head of Clinical. RN, MSN, 15 years in discharge planning.",
    ]),
    ("The Ask", [
        "Raising $6M Seed led by an institutional healthcare fund.",
        "$2.5M committed from existing investors.",
        "Priced equity round at a $28M pre-money valuation.",
    ]),
    ("Use of Funds", [
        "Engineering and product: $2.4M",
        "Sales and marketing: $2.1M",
        "Clinical and implementation: $900K",
        "G&A and reserve: $600K",
        "24 months runway to $4.2M ARR",
    ]),
]


def make_pptx(destination: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    blank = presentation.slide_layouts[5]  # title only

    for title, lines in SLIDES:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(8.8), Inches(4.6))
        frame = box.text_frame
        frame.word_wrap = True
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            for run in paragraph.runs:
                run.font.size = Pt(14)

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(destination))
    return destination


def make_pdf(destination: Path) -> Path:
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas as pdfcanvas

    page = landscape(letter)
    width, height = page
    destination.parent.mkdir(parents=True, exist_ok=True)
    canvas = pdfcanvas.Canvas(str(destination), pagesize=page)

    for title, lines in SLIDES:
        canvas.setFont("Helvetica-Bold", 26)
        canvas.drawString(0.8 * inch, height - 1.1 * inch, title)
        canvas.setFont("Helvetica", 13)
        y = height - 1.9 * inch
        for line in lines:
            for chunk in _wrap(line, 92):
                canvas.drawString(0.8 * inch, y, chunk)
                y -= 0.26 * inch
            y -= 0.06 * inch
        canvas.showPage()

    canvas.save()
    return destination


def _wrap(text: str, width: int):
    words, line, out = text.split(), "", []
    for word in words:
        candidate = (line + " " + word).strip()
        if len(candidate) > width:
            out.append(line)
            line = word
        else:
            line = candidate
    if line:
        out.append(line)
    return out


if __name__ == "__main__":
    out = ROOT / "input"
    print("wrote", make_pptx(out / "sample_deck.pptx"))
    print("wrote", make_pdf(out / "sample_deck.pdf"))
