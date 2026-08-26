"""PowerPoint pitch-deck extraction via python-pptx.

Captures slide titles, body text from every text-bearing shape (including inside
groups), tables, chart series/category labels, speaker notes, image counts and
run-level font sizes for the design review.

Legacy binary ``.ppt`` cannot be read by python-pptx; it is detected and reported
with a clear conversion instruction rather than failing obscurely.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from models.schemas import ExtractedDeck, Slide
from utils.validation import CorruptFileError, ExtractionError, sanitize_text

TITLE_MAX_CHARS = 90


def extract(path: Path, file_type: str = "pptx") -> ExtractedDeck:
    if file_type == "ppt":
        raise ExtractionError(
            "{} is a legacy binary .ppt file, which cannot be read directly.".format(path.name),
            "Open it in PowerPoint and save as .pptx, then run the review again.",
        )

    try:
        from pptx import Presentation  # type: ignore
        from pptx.exc import PackageNotFoundError  # type: ignore
    except ImportError as exc:
        raise ExtractionError(
            "python-pptx is not installed.",
            "Install it with: pip install python-pptx",
        ) from exc

    deck = ExtractedDeck(source_path=str(path), file_type="pptx", extractor="python-pptx")

    try:
        presentation = Presentation(str(path))
    except PackageNotFoundError as exc:
        raise CorruptFileError(
            "{} is not a readable PowerPoint package.".format(path.name),
            "The file may be corrupted or password-protected. Re-save it as .pptx.",
        ) from exc
    except Exception as exc:  # noqa: BLE001
        raise CorruptFileError(
            "Could not open {}: {}".format(path.name, exc),
            "Open the deck in PowerPoint and re-save it as .pptx.",
        ) from exc

    deck.metadata = _core_properties(presentation)

    for slide_index, raw_slide in enumerate(presentation.slides, start=1):
        slide = _read_slide(raw_slide, slide_index, deck)
        deck.slides.append(slide)

    if not deck.slides:
        raise ExtractionError(
            "{} contains no slides.".format(path.name),
            "Confirm the deck exported correctly.",
        )

    if not deck.has_usable_text:
        deck.warnings.append(
            "Only {} words of text across {} slides -- the deck is largely image-based, "
            "so content analysis is limited.".format(deck.total_words, deck.slide_count)
        )

    return deck


# --------------------------------------------------------------------------- #
# Slide reading
# --------------------------------------------------------------------------- #

def _read_slide(raw_slide, index: int, deck: ExtractedDeck) -> Slide:
    title = _slide_title(raw_slide)
    body_lines: List[str] = []
    tables: List[List[List[str]]] = []
    chart_labels: List[str] = []
    images = 0
    font_sizes: List[float] = []

    for shape in _walk_shapes(raw_slide.shapes):
        kind = _shape_kind(shape)

        if kind == "picture":
            images += 1
            continue

        if kind == "table":
            table = _read_table(shape)
            if table:
                tables.append(table)
            continue

        if kind == "chart":
            chart_labels.extend(_read_chart(shape, deck))
            continue

        text, sizes = _read_text_frame(shape)
        if text and text != title:
            body_lines.append(text)
        font_sizes.extend(sizes)

    slide = Slide(
        index=index,
        title=sanitize_text(title)[:TITLE_MAX_CHARS],
        text=sanitize_text("\n".join(body_lines)),
        notes=sanitize_text(_read_notes(raw_slide)),
        tables=tables,
        chart_labels=[sanitize_text(c) for c in chart_labels if c],
        image_count=images,
    )
    slide.word_count = len(slide.text.split()) + len(slide.title.split())
    slide.is_text_sparse = slide.word_count < 12 and images > 0

    stats: Dict[str, Any] = {"small_spans": sum(1 for s in font_sizes if s < 10)}
    if font_sizes:
        stats["avg_size"] = round(sum(font_sizes) / len(font_sizes), 1)
        stats["min_size"] = min(font_sizes)
        stats["max_size"] = max(font_sizes)
    deck.metadata.setdefault("font_stats", []).append(stats)

    return slide


def _walk_shapes(shapes):
    """Yield every shape, descending into groups."""
    for shape in shapes:
        if getattr(shape, "shape_type", None) is not None and _is_group(shape):
            try:
                yield from _walk_shapes(shape.shapes)
            except Exception:  # noqa: BLE001
                pass
            continue
        yield shape


def _is_group(shape) -> bool:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore

        return shape.shape_type == MSO_SHAPE_TYPE.GROUP
    except Exception:  # noqa: BLE001
        return False


def _shape_kind(shape) -> str:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return "picture"
    except Exception:  # noqa: BLE001
        pass
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_text_frame", False):
        return "text"
    return "other"


def _slide_title(raw_slide) -> str:
    try:
        placeholder = raw_slide.shapes.title
        if placeholder is not None and placeholder.has_text_frame:
            text = placeholder.text_frame.text.strip()
            if text:
                return text
    except Exception:  # noqa: BLE001
        pass
    return ""


def _read_text_frame(shape):
    """Return (text, font_sizes) for a text-bearing shape."""
    sizes: List[float] = []
    if not getattr(shape, "has_text_frame", False):
        return "", sizes
    lines: List[str] = []
    try:
        for paragraph in shape.text_frame.paragraphs:
            runs = []
            for run in paragraph.runs:
                if run.text:
                    runs.append(run.text)
                size = getattr(run.font, "size", None)
                if size is not None:
                    sizes.append(round(size.pt, 1))
            line = "".join(runs).strip()
            if line:
                lines.append(line)
    except Exception:  # noqa: BLE001
        try:
            lines = [shape.text_frame.text]
        except Exception:  # noqa: BLE001
            lines = []
    return "\n".join(l for l in lines if l.strip()), sizes


def _read_table(shape) -> List[List[str]]:
    rows: List[List[str]] = []
    try:
        for row in shape.table.rows:
            cells = [sanitize_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
    except Exception:  # noqa: BLE001
        return []
    return rows if len(rows) > 1 else []


def _read_chart(shape, deck: ExtractedDeck) -> List[str]:
    """Chart title, category names and series names -- the machine-readable parts."""
    labels: List[str] = []
    try:
        chart = shape.chart
        if getattr(chart, "has_title", False) and chart.chart_title.has_text_frame:
            labels.append(chart.chart_title.text_frame.text)
        try:
            labels.extend(str(c) for c in chart.plots[0].categories)
        except Exception:  # noqa: BLE001
            pass
        for series in chart.series:
            name = getattr(series, "name", None)
            if name:
                labels.append(str(name))
    except Exception as exc:  # noqa: BLE001
        deck.warnings.append("A chart could not be read: {}".format(exc))
    return [l for l in labels if l and l.strip()]


def _read_notes(raw_slide) -> str:
    try:
        if raw_slide.has_notes_slide:
            return raw_slide.notes_slide.notes_text_frame.text or ""
    except Exception:  # noqa: BLE001
        pass
    return ""


def _core_properties(presentation) -> Dict[str, Any]:
    meta: Dict[str, Any] = {}
    try:
        props = presentation.core_properties
        for key in ("title", "author", "subject", "category", "comments"):
            value = getattr(props, key, None)
            if value:
                meta[key] = sanitize_text(str(value))
    except Exception:  # noqa: BLE001
        pass
    return meta
